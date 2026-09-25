import os
import pandas as pd
from docx import Document
from docx.shared import Pt


def add_markdown_paragraphs(doc: Document, md_text: str):
    lines = md_text.splitlines()
    for line in lines:
        if line.startswith('#'):
            level = line.count('#', 0, line.find(' '))
            text = line.lstrip('# ').strip()
            p = doc.add_heading(text, level=min(level, 3))
        elif line.strip() == '':
            doc.add_paragraph('')
        else:
            doc.add_paragraph(line)


def add_dataframe_as_table(doc: Document, df: pd.DataFrame, title: str = None):
    if title:
        doc.add_heading(title, level=3)
    table = doc.add_table(rows=1, cols=len(df.columns))
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr_cells[i].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)


def generate_report(md_path: str, tables_dir: str, out_path: str):
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)

    # Add markdown content
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()
    add_markdown_paragraphs(doc, md_text)

    # Add CSV tables
    for fname in sorted(os.listdir(tables_dir)):
        if not fname.lower().endswith('.csv'):
            continue
        path = os.path.join(tables_dir, fname)
        try:
            df = pd.read_csv(path)
        except Exception:
            continue
        title = os.path.splitext(fname)[0].replace('_', ' ').title()
        add_dataframe_as_table(doc, df, title=title)

    doc.save(out_path)


if __name__ == '__main__':
    root = os.path.dirname(os.path.dirname(__file__))
    md_path = os.path.join(root, 'report', 'analysis_report.md')
    tables_dir = os.path.join(root, 'outputs', 'tables')
    out_path = os.path.join(root, 'report', 'analysis_report.docx')
    generate_report(md_path, tables_dir, out_path)
    print(f'Wrote {out_path}')
